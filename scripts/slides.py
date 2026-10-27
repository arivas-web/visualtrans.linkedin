#!/usr/bin/env python3
"""Genera las slides (.pptx) de entrega del contenido LinkedIn, siempre con el mismo diseño.

Es el ÚNICO formato de entrega de posts y calendarios (ver CLAUDE.md). No se escribe el .pptx
a mano ni con otra herramienta: se llama a este script.

Uso:
  # Calendario (antes de redactar): portada + "qué aprendimos" + tablas por semana
  python3 -I scripts/slides.py --titulo "Calendario noviembre 2026" \
      --calendario output/2026-11/calendario.md --retro output/2026-11/retro.md \
      --salida output/2026-11/calendario-2026-11.pptx

  # Posts redactados (mes, lote, rango o suelto): portada + "qué aprendimos" + un post por slide
  python3 -I scripts/slides.py --titulo "Posts noviembre 2026" \
      --posts output/2026-11/posts/*.md --retro output/2026-11/retro.md \
      --salida output/2026-11/posts-2026-11.pptx

Formato de los posts: bloques con cabecera entre líneas `---` (PERFIL, FECHA, PILAR, PAIN #,
TEMA, FORMATO y opcionales POR QUÉ, IMAGEN) seguidos del texto. El retro es un markdown con
secciones `## Funcionó`, `## No funcionó` y `## Qué cambiamos` (listas con guiones).
"""
import argparse
import datetime as dt
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x14, 0x23, 0x3F)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x67, 0x78)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xF1, 0xF4, 0xF8)
ORANGE = RGBColor(0xF2, 0x6B, 0x1D)
GREEN = RGBColor(0x1E, 0x8E, 0x5A)
RED = RGBColor(0xC8, 0x3A, 0x32)
FONT_T = "Cambria"
FONT_B = "Calibri"

PERFIL_COLOR = {
    "visual trans": RGBColor(0x1D, 0x5F, 0xC9),
    "emma": RGBColor(0x0E, 0x94, 0x88),
    "cecilio": RGBColor(0x6C, 0x4B, 0xD8),
    "enrique": RGBColor(0xC2, 0x45, 0x2F),
    "laura": RGBColor(0xD1, 0x86, 0x1D),
}
DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def color_perfil(perfil):
    p = perfil.lower()
    for k, c in PERFIL_COLOR.items():
        if k in p:
            return c
    return MUTED


# ---------- lectura ----------

def leer_posts(rutas):
    posts = []
    for ruta in rutas:
        lineas = Path(ruta).read_text(encoding="utf-8").splitlines()
        i = 0
        while i < len(lineas):
            if lineas[i].strip() == "---" and i + 1 < len(lineas) and lineas[i + 1].startswith("PERFIL:"):
                cab, j = {}, i + 1
                while j < len(lineas) and lineas[j].strip() != "---":
                    if ":" in lineas[j]:
                        k, v = lineas[j].split(":", 1)
                        cab[k.strip().upper()] = v.strip()
                    j += 1
                j += 1
                cuerpo = []
                while j < len(lineas):
                    if lineas[j].strip() == "---" and (j + 1 >= len(lineas) or lineas[j + 1].startswith("PERFIL:")
                                                        or lineas[j + 1].startswith("## ")):
                        break
                    if lineas[j].startswith("## Notas de redacción"):
                        break
                    cuerpo.append(lineas[j])
                    j += 1
                cab["TEXTO"] = "\n".join(cuerpo).strip()
                if not cab.get("PILAR"):
                    tema = cab.get("TEMA", "").lower()
                    cab["PILAR"] = "Pains" if tema.startswith("pain") else "Noticias" if tema else ""
                posts.append(cab)
                i = j
            else:
                i += 1
    posts.sort(key=lambda p: (clave_fecha(p.get("FECHA", "")), p.get("PERFIL", "")))
    return posts


def clave_fecha(f):
    m = re.search(r"(\d{1,2})/(\d{1,2})", f or "")
    if not m:
        return (99, 99)
    return (int(m.group(2)), int(m.group(1)))


def leer_calendario(ruta):
    filas, cab = [], None
    for linea in Path(ruta).read_text(encoding="utf-8").splitlines():
        if not linea.startswith("|"):
            if cab and filas:
                break
            continue
        celdas = [c.strip() for c in linea.strip().strip("|").split("|")]
        if cab is None:
            cab = celdas
        elif set("".join(celdas)) <= set("-: "):
            continue
        else:
            filas.append(dict(zip(cab, celdas)))
    return filas


def leer_retro(ruta):
    if not ruta or not Path(ruta).exists():
        return None
    secciones, actual = {}, None
    for linea in Path(ruta).read_text(encoding="utf-8").splitlines():
        if linea.startswith("## "):
            actual = linea[3:].strip().lower()
            secciones[actual] = []
        elif actual and re.match(r"^\s*[-*] ", linea):
            secciones[actual].append(re.sub(r"^\s*[-*] ", "", linea).replace("**", "").strip())
    retro = {"si": [], "no": [], "cambia": []}
    for nombre, items in secciones.items():
        if nombre.startswith("no funcion"):
            retro["no"] += items
        elif nombre.startswith("funcion"):
            retro["si"] += items
        elif nombre.startswith("qué cambia") or nombre.startswith("que cambia"):
            retro["cambia"] += items
    return retro


def col(fila, *claves):
    for k, v in fila.items():
        if any(c in k.lower() for c in claves):
            return v
    return ""


# ---------- dibujo ----------

def caja(slide, x, y, w, h, fill=None, forma=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def texto(slide, x, y, w, h, contenido, size=14, color=INK, bold=False, font=FONT_B,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    parrafos = contenido if isinstance(contenido, list) else str(contenido).split("\n")
    for i, p in enumerate(parrafos):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = p
        run.font.size = Pt(size)
        run.font.name = font
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return tb


def chip(slide, x, y, etiqueta, fill, size=12):
    w = 0.25 + 0.085 * len(etiqueta) * size / 12
    s = caja(slide, x, y, w, 0.34, fill, MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = etiqueta
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT_B
    r.font.color.rgb = PAPER
    return w


def portada(prs, titulo, subtitulo, resumen):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    caja(s, 0, 0, 13.333, 7.5, NAVY)
    texto(s, 0.8, 0.7, 6, 0.4, "VISUAL TRANS · LINKEDIN", 12, RGBColor(0x9F, 0xB3, 0xD6), bold=True)
    texto(s, 0.8, 2.1, 11.5, 1.6, titulo, 44, PAPER, bold=True, font=FONT_T)
    texto(s, 0.8, 3.75, 11.5, 0.6, subtitulo, 18, RGBColor(0xC9, 0xD6, 0xEC))
    x = 0.8
    for num, etiqueta in resumen:
        texto(s, x, 5.2, 2.6, 0.8, str(num), 40, ORANGE, bold=True, font=FONT_T)
        texto(s, x, 6.0, 2.6, 0.5, etiqueta, 13, RGBColor(0xC9, 0xD6, 0xEC))
        x += 2.9


def slide_retro(prs, retro):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    texto(s, 0.6, 0.45, 12, 0.7, "Qué aprendimos y qué cambia", 32, NAVY, bold=True, font=FONT_T)
    cols = [("Funcionó", retro["si"], GREEN), ("No funcionó (no se repite)", retro["no"], RED),
            ("Qué cambiamos", retro["cambia"], ORANGE)]
    for i, (tit, items, c) in enumerate(cols):
        x = 0.6 + i * 4.1
        caja(s, x, 1.45, 3.85, 5.55, SOFT, MSO_SHAPE.ROUNDED_RECTANGLE)
        dot = caja(s, x + 0.25, 1.7, 0.32, 0.32, c, MSO_SHAPE.OVAL)
        texto(s, x + 0.7, 1.68, 3.0, 0.4, tit, 16, c, bold=True)
        items = items[:6] or ["(sin datos todavía)"]
        largo = sum(len(t) for t in items)
        size = 14 if largo < 420 else 12 if largo < 650 else 11
        texto(s, x + 0.25, 2.3, 3.35, 4.55, ["• " + t for t in items], size, INK)


def slides_tabla(prs, titulo, filas, columnas, anchos, por_slide=11):
    for n in range(0, len(filas), por_slide):
        trozo = filas[n:n + por_slide]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        texto(s, 0.6, 0.45, 12, 0.7, titulo, 28, NAVY, bold=True, font=FONT_T)
        tabla = s.shapes.add_table(len(trozo) + 1, len(columnas), Inches(0.6), Inches(1.3),
                                   Inches(12.1), Inches(0.4 * (len(trozo) + 1))).table
        for j, (nombre, ancho) in enumerate(zip(columnas, anchos)):
            tabla.columns[j].width = Inches(ancho)
            celda = tabla.cell(0, j)
            celda.fill.solid()
            celda.fill.fore_color.rgb = NAVY
            _celda(celda, nombre, PAPER, True)
        for i, fila in enumerate(trozo, start=1):
            for j, valor in enumerate(fila):
                celda = tabla.cell(i, j)
                celda.fill.solid()
                celda.fill.fore_color.rgb = SOFT if i % 2 else PAPER
                _celda(celda, valor, color_perfil(valor) if j == 1 else INK, j == 1)


def _celda(celda, valor, color, bold):
    tf = celda.text_frame
    tf.word_wrap = True
    celda.margin_left = celda.margin_right = Inches(0.06)
    celda.margin_top = celda.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = valor
    r.font.size = Pt(10)
    r.font.name = FONT_B
    r.font.bold = bold
    r.font.color.rgb = color


def slide_post(prs, p, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    perfil = p.get("PERFIL", "")
    c = color_perfil(perfil)
    caja(s, 0, 0, 4.2, 7.5, SOFT)
    texto(s, 0.5, 0.5, 3.4, 0.3, f"{n} / {total}", 11, MUTED)
    texto(s, 0.5, 0.85, 3.4, 0.7, p.get("FECHA", "sin fecha"), 26, NAVY, bold=True, font=FONT_T)
    chip(s, 0.5, 1.7, perfil or "—", c, 13)
    filas = [("Pilar", p.get("PILAR", "")), ("Pain", p.get("PAIN #", "")), ("Formato", p.get("FORMATO", "")),
             ("Tema", p.get("TEMA", "")), ("Por qué / qué probamos", p.get("POR QUÉ", "") or p.get("NOTA", ""))]
    y = 2.35
    for etiqueta, valor in filas:
        if not valor or valor == "—":
            continue
        texto(s, 0.5, y, 3.3, 0.25, etiqueta.upper(), 10, MUTED, bold=True)
        alto = 0.3 + 0.21 * (len(valor) // 34)
        texto(s, 0.5, y + 0.27, 3.3, alto, valor, 12, INK)
        y += 0.42 + alto
        if y > 6.6:
            break
    img = p.get("IMAGEN", "")
    if img.startswith("http"):
        tb = texto(s, 0.5, 6.85, 3.3, 0.3, "Ver imagen →", 12, ORANGE, bold=True)
        tb.text_frame.paragraphs[0].runs[0].hyperlink.address = img
    cuerpo = p.get("TEXTO", "")
    largo = len(cuerpo)
    size = 15 if largo < 700 else 13 if largo < 1100 else 12 if largo < 1500 else 11 if largo < 1900 else 10
    parrafos = [l for l in cuerpo.split("\n")]
    texto(s, 4.8, 0.55, 7.9, 6.5, parrafos, size, INK)
    texto(s, 4.8, 7.05, 7.9, 0.3, f"{largo} caracteres", 10, MUTED, align=PP_ALIGN.RIGHT)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--titulo", required=True)
    ap.add_argument("--subtitulo", default="")
    ap.add_argument("--posts", nargs="*", default=[])
    ap.add_argument("--calendario")
    ap.add_argument("--retro")
    ap.add_argument("--salida", required=True)
    a = ap.parse_args()
    if not a.posts and not a.calendario:
        sys.exit("Indica --posts o --calendario")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    retro = leer_retro(a.retro)
    hoy = dt.date.today().strftime("%d/%m/%Y")

    if a.posts:
        posts = leer_posts(a.posts)
        perfiles = {p.get("PERFIL") for p in posts}
        resumen = [(len(posts), "posts"), (len(perfiles), "perfiles"),
                   (sum("pain" in p.get("PILAR", "").lower() for p in posts), "pains"),
                   (sum("noticia" in p.get("PILAR", "").lower() for p in posts), "noticias")]
        portada(prs, a.titulo, a.subtitulo or f"Para revisar · generado el {hoy}", resumen)
        if retro:
            slide_retro(prs, retro)
        if len(posts) > 3:
            filas = [[p.get("FECHA", ""), p.get("PERFIL", ""), p.get("PILAR", ""), p.get("FORMATO", ""),
                      p.get("TEMA", "")] for p in posts]
            slides_tabla(prs, "De un vistazo", filas, ["Fecha", "Perfil", "Pilar", "Formato", "Tema"],
                         [1.4, 1.9, 1.3, 2.2, 5.3])
        for n, p in enumerate(posts, 1):
            slide_post(prs, p, n, len(posts))
    else:
        filas = leer_calendario(a.calendario)
        perfiles = {col(f, "perfil") for f in filas}
        resumen = [(len(filas), "posts"), (len(perfiles), "perfiles"),
                   (sum("pain" in col(f, "pilar").lower() for f in filas), "pains"),
                   (sum("noticia" in col(f, "pilar").lower() for f in filas), "noticias")]
        portada(prs, a.titulo, a.subtitulo or f"Pendiente de validar · generado el {hoy}", resumen)
        if retro:
            slide_retro(prs, retro)
        semanas = {}
        for f in filas:
            semanas.setdefault(col(f, "semana") or "—", []).append(f)
        for sem, fs in semanas.items():
            datos = [[f"{col(f, 'día', 'dia')[:3]} {col(f, 'fecha')}", col(f, "perfil"), col(f, "pilar"),
                      col(f, "formato"), col(f, "tema"), col(f, "por qué", "porque")] for f in fs]
            slides_tabla(prs, f"Semana {sem}", datos, ["Fecha", "Perfil", "Pilar", "Formato", "Tema / enfoque",
                                                       "Por qué"], [1.1, 1.6, 1.0, 1.6, 4.4, 2.4], por_slide=9)

    Path(a.salida).parent.mkdir(parents=True, exist_ok=True)
    prs.save(a.salida)
    print(a.salida)


if __name__ == "__main__":
    main()
